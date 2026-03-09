"""엑셀/CSV 업로드 비동기 작업 서비스

인메모리 싱글턴으로 upload 작업 상태를 추적한다.
SyncJobService 패턴을 재사용하되, 파일 파싱 데이터는 state에 저장하지 않는다.
(500MB 컨테이너 메모리 제약 대응)
"""

from __future__ import annotations

import asyncio
import csv
import gc
import io
import logging
import re
import time
from dataclasses import dataclass

from schemas.upload import UploadJobStatusResponse, UploadResultResponse, UploadRowError
from services.base_job_service import BaseJobService, BaseJobState

logger = logging.getLogger(__name__)

# --- 컬럼 매핑 상수 (import_from_csv.py에서 추출) ---

STATUS_MAP = {
    "정상": "normal",
    "블라인드": "blind",
    "삭제": "deleted",
}

_HTML_TAG_RE = re.compile(r"<[^>]+>")

PIPELINE_CHUNK = 500

# 엑셀/CSV 한글 → 영문 컬럼 매핑
COLUMN_MAP = {
    "리뷰번호": "review_id",
    "지점번호": "branch_id",
    "리뷰내용": "content",
    "예약_지점명": "branch_name",
    "예약_업체명": "company_name",
    "지점평점(친절/편의성)": "rating_service",
    "차량평점": "rating_car",
    "인수/반납편의성": "rating_convenience",
    "도움돼요수": "helpful_count",
    "등록일시": "review_date",
    "리뷰상태": "status",
    "차량모델": "car_type",
    "렌트타입": "rent_type",
}

# 필수 컬럼 (한글 또는 영문)
REQUIRED_COLUMNS_KR = {"리뷰번호", "지점번호"}
REQUIRED_COLUMNS_EN = {"review_id", "branch_id"}


def _strip_null(text: str) -> str:
    """PostgreSQL이 거부하는 null byte 제거"""
    return text.replace("\x00", "") if text else text


def _strip_html(text: str) -> str:
    """HTML 태그 + null byte 제거 후 공백 정리"""
    if not text:
        return ""
    cleaned = _HTML_TAG_RE.sub(" ", text)
    cleaned = cleaned.replace("\x00", "")
    return re.sub(r"\s{2,}", " ", cleaned).strip()


ParseResult = tuple[list[dict], list[dict], list[UploadRowError]]


def parse_rows(raw_dicts: list[dict]) -> ParseResult:
    """딕셔너리 리스트를 (raw_rows, mapped_rows, errors) 튜플로 변환.

    raw_rows: 한국어 키 그대로 (upsert_batch 전달용)
    mapped_rows: 영어 키 변환 (UnifiedPipeline.run() 전달용)
    errors: 행별 검증 오류 목록
    """
    raw_rows: list[dict] = []
    mapped_rows: list[dict] = []
    errors: list[UploadRowError] = []

    for idx, row in enumerate(raw_dicts, start=2):  # 헤더가 1행이므로 데이터는 2행부터
        # 한글 키인지 영문 키인지 판단
        has_kr = "리뷰번호" in row
        has_en = "review_id" in row

        if not has_kr and not has_en:
            msg = "리뷰번호(review_id) 컬럼 누락"
            errors.append(UploadRowError(row=idx, error=msg))
            continue

        if has_kr:
            review_id = (row.get("리뷰번호") or "").strip()
            branch_id = (row.get("지점번호") or "").strip()
        else:
            review_id = (row.get("review_id") or "").strip()
            branch_id = (row.get("branch_id") or "").strip()

        if not review_id:
            errors.append(UploadRowError(row=idx, error="review_id 값 누락"))
            continue
        if not branch_id:
            errors.append(UploadRowError(row=idx, error="branch_id 값 누락"))
            continue

        # raw_rows: null byte만 제거
        raw_row = {
            k: _strip_null(v) if isinstance(v, str) else v
            for k, v in row.items()
        }
        raw_row["is_new"] = True
        raw_rows.append(raw_row)

        # mapped_rows: 영어 키 변환
        if has_kr:
            status_kr = (row.get("리뷰상태") or "").strip()
            status_en = STATUS_MAP.get(status_kr, status_kr)
            content_clean = _strip_html(row.get("리뷰내용") or "")

            mapped_rows.append({
                "review_id": review_id,
                "branch_id": branch_id,
                "content": content_clean,
                "branch_name": (row.get("예약_지점명") or "").strip(),
                "company_name": (row.get("예약_업체명") or "").strip(),
                "rating_service": (row.get("지점평점(친절/편의성)") or "").strip(),
                "rating_car": (row.get("차량평점") or "").strip(),
                "rating_convenience": (row.get("인수/반납편의성") or "").strip(),
                "helpful_count": (row.get("도움돼요수") or "0").strip(),
                "review_date": (row.get("등록일시") or "").strip(),
                "status": status_en,
                "car_type": (row.get("차량모델") or "").strip(),
                "rent_type": (row.get("렌트타입") or "").strip(),
            })
        else:
            # 이미 영문 키인 경우 content만 정리
            content_clean = _strip_html(row.get("content") or "")
            mapped_row = {**row, "content": content_clean}
            mapped_rows.append(mapped_row)

    return raw_rows, mapped_rows, errors


def parse_csv_content(content: bytes) -> list[dict]:
    """CSV 바이트 → 딕셔너리 리스트"""
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return list(reader)


def parse_excel_content(content: bytes) -> list[dict]:
    """Excel 바이트 → 딕셔너리 리스트"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return []

    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        wb.close()
        return []

    result: list[dict] = []
    for row_values in rows_iter:
        row_dict = {}
        for header, value in zip(headers, row_values, strict=False):
            if header:
                row_dict[header] = str(value).strip() if value is not None else ""
        if any(row_dict.values()):  # 빈 행 스킵
            result.append(row_dict)

    wb.close()
    return result


def validate_columns(raw_dicts: list[dict]) -> str | None:
    """필수 컬럼 존재 여부 검증. 오류 시 메시지 반환, 정상이면 None."""
    if not raw_dicts:
        return "파일에 데이터가 없습니다"

    columns = set(raw_dicts[0].keys())
    has_kr = REQUIRED_COLUMNS_KR.issubset(columns)
    has_en = REQUIRED_COLUMNS_EN.issubset(columns)

    if not has_kr and not has_en:
        missing_kr = REQUIRED_COLUMNS_KR - columns
        missing_en = REQUIRED_COLUMNS_EN - columns
        return (
            f"필수 컬럼 누락. "
            f"한글({', '.join(sorted(missing_kr))}) 또는 "
            f"영문({', '.join(sorted(missing_en))}) 컬럼이 필요합니다. "
            f"발견된 컬럼: {', '.join(sorted(columns))}"
        )

    return None


# --- Job State ---


@dataclass
class UploadJobState(BaseJobState):
    """인메모리 업로드 작업 상태 (파싱 데이터는 저장하지 않음)"""

    total_rows: int = 0
    result: UploadResultResponse | None = None


class UploadJobService(BaseJobService[UploadJobState, UploadJobStatusResponse]):
    """엑셀/CSV 업로드 비동기 작업 관리 (인메모리 싱글턴)"""

    async def submit_job(
        self,
        raw_rows: list[dict],
        mapped_rows: list[dict],
        total_rows: int,
        filename: str = "",
    ) -> UploadJobStatusResponse:
        """비동기 업로드 작업 제출.

        이미 활성 작업이 있으면 기존 job_id를 반환한다 (단일 활성 업로드 제한).
        raw_rows/mapped_rows는 background task의 로컬 변수로만 유지된다 (GC 대응).
        """
        active = self._find_active_job()
        if active:
            return self._to_response(active)

        self._prune_old_jobs()

        job_id = self._generate_job_id()
        state = UploadJobState(
            job_id=job_id,
            total_rows=total_rows,
            message=f"파일 업로드 완료, 처리 대기 중 ({filename})",
        )
        self._jobs[job_id] = state

        # raw_rows/mapped_rows를 task 인자로 전달 (state에 저장하지 않음)
        task = asyncio.create_task(self._run_job(state, raw_rows, mapped_rows))
        state.task = task

        logger.info(
            "업로드 작업 제출: job_id=%s, rows=%d, file=%s",
            job_id, total_rows, filename,
        )
        return self._to_response(state)

    async def _run_job(
        self,
        state: UploadJobState,
        raw_rows: list[dict],
        mapped_rows: list[dict],
    ) -> None:
        """백그라운드에서 업로드 파이프라인 실행"""
        start = time.time()
        try:
            state.status = "processing"
            state.progress = 5
            state.message = "DB 저장 준비 중"

            from domain.pipeline.unified_pipeline import UnifiedPipeline
            from repository.database import get_session_factory
            from repository.review_repository import BranchReviewRepository

            # --- Phase 1: upsert_batch (별도 세션, 완료 후 커밋) ---
            factory = get_session_factory()
            upserted_count = 0
            async with factory() as session:
                try:
                    review_repo = BranchReviewRepository(session)
                    upserted_count = await review_repo.upsert_batch(raw_rows)

                    # 데드락 방지: UnifiedPipeline이 별도 세션으로 같은 행을 UPDATE하므로
                    # 행 잠금을 해제해야 함 (sync_service.py:122-124 패턴)
                    await session.commit()
                    logger.info(
                        "업로드 upsert 완료: %d건 (job=%s)",
                        upserted_count, state.job_id,
                    )
                except Exception:
                    await session.rollback()
                    raise

            state.progress = 30
            state.message = f"{upserted_count}건 저장 완료, 분석 시작"

            # --- Phase 2: UnifiedPipeline 청크 실행 (파이프라인 자체 세션 사용) ---
            pipeline = UnifiedPipeline()
            total = len(mapped_rows)
            total_processed = 0

            chunk_count = (total + PIPELINE_CHUNK - 1) // PIPELINE_CHUNK
            for i in range(0, total, PIPELINE_CHUNK):
                chunk = mapped_rows[i : i + PIPELINE_CHUNK]
                chunk_no = i // PIPELINE_CHUNK + 1

                result = await pipeline.run(chunk)
                total_processed += result.processed_reviews

                # 진행률: 30% (upsert) ~ 95% (pipeline)
                pct = 30 + int(min(i + PIPELINE_CHUNK, total) / total * 65)
                state.progress = pct
                state.message = (
                    f"분석 중 ({chunk_no}/{chunk_count} 청크, "
                    f"{total_processed}/{total}건 처리)"
                )
                gc.collect()

            # --- 완료 ---
            duration = time.time() - start
            failed_count = total - total_processed

            state.status = "completed"
            state.progress = 100
            state.message = f"{upserted_count}건 저장 + {total_processed}건 분석 완료"
            state.result = UploadResultResponse(
                success=True,
                uploaded_count=state.total_rows,
                upserted_count=upserted_count,
                processed_count=total_processed,
                failed_count=failed_count,
                duration_seconds=round(duration, 1),
            )
            logger.info(
                "업로드 작업 완료: job=%s, upserted=%d, processed=%d, failed=%d, %.1fs",
                state.job_id, upserted_count, total_processed, failed_count, duration,
            )

        except asyncio.CancelledError:
            state.status = "failed"
            state.error = "작업이 취소되었습니다"
            logger.info("업로드 작업 취소됨: %s", state.job_id)
        except Exception as e:
            state.status = "failed"
            state.error = str(e)
            state.message = "업로드 처리 중 오류 발생"
            logger.exception("업로드 작업 실패: %s", state.job_id)

    def _to_response(self, state: UploadJobState) -> UploadJobStatusResponse:
        return UploadJobStatusResponse(
            job_id=state.job_id,
            status=state.status,
            progress=state.progress,
            message=state.message,
            total_rows=state.total_rows,
            error=state.error,
            result=state.result,
            created_at=state.created_at,
        )
