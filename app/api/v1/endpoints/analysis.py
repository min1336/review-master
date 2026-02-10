"""
리뷰 분석 페이지 API

Router: /api/analysis
담당: HTTP 요청/응답 처리만
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from urllib.parse import quote

from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from schemas.common import api_response
from services.analysis_service import AnalysisService

from .deps import get_analysis_service

logger = logging.getLogger(__name__)

router = APIRouter(tags=["analysis"])


@router.get("/filters")
async def get_filter_options(
    service: AnalysisService = Depends(get_analysis_service),
) -> dict[str, Any]:
    """
    필터 옵션 조회 (지역, 업체명, 지점 목록)

    Returns:
        필터 옵션 데이터
            - regions: 지역 목록
            - companies: 업체명 목록
            - branches: 지점 목록 [{branch_id, branch_name}, ...]
    """
    try:
        result = await service.get_filter_options()
        return api_response(result.to_dict())

    except Exception as e:
        logger.error(f"Failed to get filter options: {e}")
        raise HTTPException(
            status_code=500,
            detail="필터 옵션을 불러오는데 실패했습니다.",
        )


@router.get("/reviews")
async def get_reviews(
    regions: list[str] | None = Query(None, description="지역 필터 목록"),
    companies: list[str] | None = Query(None, description="업체명 필터 목록"),
    branch_ids: list[int] | None = Query(None, description="지점 ID 필터 목록"),
    sentiment: str | None = Query(
        None,
        pattern="^(positive|negative|neutral)$",
        description="감정 필터 (positive, negative, neutral)",
    ),
    date_from: str | None = Query(
        None,
        pattern=r"^\d{4}-\d{2}-\d{2}$",
        description="시작일 (YYYY-MM-DD)",
    ),
    date_to: str | None = Query(
        None,
        pattern=r"^\d{4}-\d{2}-\d{2}$",
        description="종료일 (YYYY-MM-DD)",
    ),
    sort_by: str = Query(
        "latest",
        pattern="^(latest|rating_low)$",
        description="정렬 기준 (latest, rating_low)",
    ),
    limit: int = Query(20, ge=1, le=100, description="조회 개수 (기본 20, 최대 100)"),
    offset: int = Query(0, ge=0, description="페이징 오프셋"),
    is_new: bool | None = Query(None, description="신규 리뷰 필터 (true: 신규만, false: 읽은 것만)"),
    service: AnalysisService = Depends(get_analysis_service),
) -> dict[str, Any]:
    """
    필터링된 리뷰 조회

    Returns:
        리뷰 목록과 전체 개수
            - reviews: 리뷰 목록
            - total: 전체 개수
    """
    try:
        result = await service.get_filtered_reviews(
            regions=regions,
            companies=companies,
            branch_ids=branch_ids,
            sentiment=sentiment,
            date_from=date_from,
            date_to=date_to,
            sort_by=sort_by,
            limit=limit,
            offset=offset,
            is_new=is_new,
        )
        return api_response(result.to_dict())

    except Exception as e:
        logger.error(f"Failed to get filtered reviews: {e}")
        raise HTTPException(
            status_code=500,
            detail="리뷰를 불러오는데 실패했습니다.",
        )


@router.get("/reviews/export")
async def export_reviews_to_excel(
    regions: list[str] | None = Query(None, description="지역 필터 목록"),
    companies: list[str] | None = Query(None, description="업체명 필터 목록"),
    branch_ids: list[int] | None = Query(None, description="지점 ID 필터 목록"),
    sentiment: str | None = Query(
        None,
        pattern="^(positive|negative|neutral)$",
        description="감정 필터 (positive, negative, neutral)",
    ),
    date_from: str | None = Query(
        None,
        pattern=r"^\d{4}-\d{2}-\d{2}$",
        description="시작일 (YYYY-MM-DD)",
    ),
    date_to: str | None = Query(
        None,
        pattern=r"^\d{4}-\d{2}-\d{2}$",
        description="종료일 (YYYY-MM-DD)",
    ),
    sort_by: str = Query(
        "latest",
        pattern="^(latest|rating_low)$",
        description="정렬 기준 (latest, rating_low)",
    ),
    service: AnalysisService = Depends(get_analysis_service),
) -> StreamingResponse:
    """
    필터링된 리뷰를 엑셀 파일로 내보내기

    Returns:
        StreamingResponse: 엑셀 파일 다운로드
    """
    try:
        # 최대 10000개까지 내보내기
        result = await service.get_filtered_reviews(
            regions=regions,
            companies=companies,
            branch_ids=branch_ids,
            sentiment=sentiment,
            date_from=date_from,
            date_to=date_to,
            sort_by=sort_by,
            limit=10000,
            offset=0,
        )

        # 엑셀 워크북 생성
        wb = Workbook()
        ws = wb.active
        ws.title = "리뷰 목록"

        # 스타일 정의
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # 헤더 작성
        headers = [
            "번호",
            "리뷰ID",
            "업체명",
            "지점명",
            "리뷰 날짜",
            "감정",
            "서비스 평점",
            "차량 평점",
            "편의성 평점",
            "리뷰 내용",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 감정 한글 변환
        sentiment_map = {
            "positive": "긍정",
            "negative": "부정",
            "neutral": "중립",
        }

        # 데이터 작성
        for row_idx, review in enumerate(result.reviews, 2):
            ws.cell(row=row_idx, column=1, value=row_idx - 1).border = thin_border
            ws.cell(row=row_idx, column=2, value=review.review_id).border = thin_border
            ws.cell(row=row_idx, column=3, value=review.company_name).border = thin_border
            ws.cell(row=row_idx, column=4, value=review.branch_name).border = thin_border

            # 날짜 포맷
            date_str = ""
            if review.review_date:
                try:
                    date_obj = datetime.fromisoformat(review.review_date.replace("Z", "+00:00"))
                    date_str = date_obj.strftime("%Y-%m-%d")
                except (ValueError, AttributeError):
                    date_str = review.review_date
            ws.cell(row=row_idx, column=5, value=date_str).border = thin_border

            ws.cell(row=row_idx, column=6, value=sentiment_map.get(review.sentiment, "-")).border = thin_border
            ws.cell(row=row_idx, column=7, value=review.rating_service).border = thin_border
            ws.cell(row=row_idx, column=8, value=review.rating_car).border = thin_border
            ws.cell(row=row_idx, column=9, value=review.rating_convenience).border = thin_border

            content_cell = ws.cell(row=row_idx, column=10, value=review.content)
            content_cell.border = thin_border
            content_cell.alignment = Alignment(wrap_text=True)

        # 열 너비 조정
        column_widths = [8, 12, 15, 20, 12, 8, 12, 12, 12, 60]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

        # 엑셀 파일을 메모리에 저장
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        # 파일명 생성: 업체_지점_감정_정렬_기간_자료
        # 업체명 결정
        if companies and len(companies) == 1:
            company_part = companies[0]
        elif companies and len(companies) > 1:
            company_part = f"{companies[0]}외{len(companies)-1}"
        else:
            company_part = "전체업체"

        # 지점명 결정
        if branch_ids and len(branch_ids) == 1:
            # 단일 지점 선택 시 지점명 찾기
            branch_name = None
            for review in result.reviews:
                if review.branch_name:
                    branch_name = review.branch_name
                    break
            branch_part = branch_name if branch_name else "지점"
        elif branch_ids and len(branch_ids) > 1:
            branch_part = f"{len(branch_ids)}개지점"
        else:
            branch_part = "전체지점"

        # 감정 결정
        sentiment_names = {"positive": "긍정", "negative": "부정", "neutral": "중립"}
        sentiment_part = sentiment_names.get(sentiment, "전체감정") if sentiment else "전체감정"

        # 정렬 결정
        sort_names = {"latest": "최신순", "rating_low": "낮은평점순"}
        sort_part = sort_names.get(sort_by, "최신순")

        # 기간 결정
        if date_from and date_to:
            period_part = f"{date_from}~{date_to}"
        elif date_from:
            period_part = f"{date_from}~"
        elif date_to:
            period_part = f"~{date_to}"
        else:
            period_part = "전체기간"

        filename = f"{company_part}_{branch_part}_{sentiment_part}_{sort_part}_{period_part}_자료.xlsx"
        encoded_filename = quote(filename)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
            },
        )

    except Exception as e:
        logger.error(f"Failed to export reviews: {e}")
        raise HTTPException(
            status_code=500,
            detail="리뷰 내보내기에 실패했습니다.",
        )
